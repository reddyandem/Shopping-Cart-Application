import os
from docx import Document
from docx.shared import Pt

# Change this path to your project directory
project_dir = 'C:/Users/prudh/Desktop/ecommerce-main'  
output_docx = 'Project_Code.docx'

document = Document()
document.add_heading('Project Code Overview', 0)

# Optionally, set a default font size for the document
style = document.styles['Normal']
font = style.font
font.name = 'Consolas'
font.size = Pt(10)

for root, dirs, files in os.walk(project_dir):
    for file in files:
        # Filter file types if needed, e.g., only Java, HTML, etc.
        if file.endswith(('.java', '.html', '.xml', '.properties')):
            file_path = os.path.join(root, file)
            document.add_heading(file_path, level=2)
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    code = f.read()
            except Exception as e:
                code = f"Error reading file: {e}"
            # Add code in a paragraph with a monospaced font style
            p = document.add_paragraph(code)
            p.style = document.styles['Normal']
            document.add_page_break()

document.save(output_docx)
print(f"Document saved as {output_docx}")
